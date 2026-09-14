"""Build the manuscript and figures using the bundled artifact runtime."""
from pathlib import Path
import json, re, hashlib
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
from pdf2image import convert_from_path

BASE=Path(__file__).resolve().parent
ROOT=BASE.parents[1]
TMP=ROOT/'tmp/paper_20260914'
TMP.mkdir(parents=True,exist_ok=True)
DATA=json.loads((BASE/'verified_evidence.json').read_text())
STEM='Vacant_可執行驗收與可究責交付'
TABLES={}
def table(key,title,header,rows,note,widths=None):
    TABLES[key]=(title,header,rows,note,widths)

table('literature','表 1　相關研究與本文的比較位置',
 ['研究脈絡','代表工作','本文承接內容','尚不能宣稱'],[
 ['測試與題庫','Barr；HumanEval；EvalPlus；LiveCodeBench [1–4]','以可執行測資評估候選','測試等同完整需求；子集等同官方全庫'],
 ['聚合與修復','Self-consistency；Self-Debugging；Reflexion；Olausson [5–8]','取樣、聚合與失敗回饋的成本比較','迴圈首創；所有條件下修復勝過重抽'],
 ['訊息與信譽','Kim；DRF；可信度評分；A-Trust [9–12]','辨識相關錯誤與裁決訊號來源','已擊敗三篇方法；多數決必然無效'],
 ['可究責紀錄','PeerReview；CT [13–14]','簽署、連續性、外部核對的原則','繼承全部系統保證；鏈通過即內容真實'],
 ['構念與方法','Körber；Nosek [15–16]','區分心理構念、事前預測與事後解釋','已提升人類信任；已通過外部事前評審']],
 '本表比較設計與證據種類，不是不同論文之間的效能排名。',[2.1,4.1,4.7,5.3])

table('arms','表 2　各實驗臂的決策、預算與出現階段',
 ['臂','候選與交付規則','最大模型呼叫','出現階段','所辨識的問題'],[
 ['OFF','一次生成後交付；隱藏測試只作事後評分','1','R460、R460R、R529','單次生成基準'],
 ['CONFORM','依序生成；第一份可見通過即交付；耗盡則拒交','5；可早停','R460、R460R、R529','驗收與額外取樣的合併收益'],
 ['OFF5','生成五份；依可見執行行為簽名多數選擇','5；固定','R460、R460R','程式行為聚合基準'],
 ['EQ5','同一組五份候選供驗收選擇與多數決使用','5；固定','r446、r448、r449b、r449c','固定生成成本後的選擇規則'],
 ['H-PI','初稿後最多四次修訂；保留較完整脈絡','5；可早停','R460、R460R','較簡單的回饋修訂'],
 ['H-OC','計畫一次、初稿一次；最多三次修訂','5；可早停','R460、R460R','包含計畫與診斷的組合'],
 ['H-MIX','初稿後回饋修訂；診斷、驗證提示及重複失敗停止','5；可早停','R460、R460R、R529','主要受評估修訂流程']],
 'H 臂另有輪間 token 與時間預算。相同呼叫額度不等於相同實際成本。EQ5 不在 R460 與 R460R 的六臂之內，各階段不是同場七臂比較。[E1–E3、E8]',[2.0,5.2,2.4,3.0,3.6])

table('eq5','表 3　EQ5 的同候選選擇規則比較',
 ['階段／題庫','驗收成功','多數決成功','b／c','Δ pp','p 原始'],[
 ['r446 MBPP+','280/371','265/371','24/9','+4.04','.0135'],
 ['r448 MBPP+','286/371','273/371','21/8','+3.50','.0241'],
 ['r449b LCB v2','85/120','75/120','15/5','+8.33','.0414'],
 ['r449c LCB v3','157/189','149/189','13/5','+4.23','.0963']],
 '各階段單獨報告，MBPP+ 兩次重用題目；不合併 n 或另作合併顯著性宣稱。來源 [E8]。',[4.3,2.5,2.6,1.8,2.2,2.8])

labels=['R460','r1','r2','r3','r4','r5']
rows=[]
for a in ['OFF','CONFORM','OFF5','HPI','HOC','HMIX']:
    cells=[{'HPI':'H-PI','HOC':'H-OC','HMIX':'H-MIX'}.get(a,a)]
    for k in labels:
        v=DATA[k]['arms'][a]
        cells.append(f"{v['correct']}/{v['n']}\n{100*v['correct']/v['n']:.2f}%")
    rows.append(cells)
table('rep_rates','表 4　初次實驗與五次複製的交付成功',
 ['臂']+labels,rows,
 '每格依序為正確交付件數／有效題數與百分比。r5 分母依臂不同；各次配對另取共同題目。來源 [E2、E4、E12]。',[2.6]+[2.27]*6)

states=['EFFECTIVE','INCONCLUSIVE','INCONCLUSIVE','RULED_OUT','INCONCLUSIVE','INCONCLUSIVE']
rows=[]
for k,state in zip(labels,states):
    v=DATA[k]['pairs']['HMIX_vs_CONFORM'];lo,hi=v['conditional_ci_pp']
    rows.append([k,str(v['n']),f"{v['b']}/{v['c']}",f"{v['delta_pp']:+.2f}",f'[{lo:.2f}, {hi:.2f}]',f"{v['p_holm']:.4f}",state])
table('rep_effects','表 5　H-MIX 相對 CONFORM 的配對效果',
 ['階段','共同 n','b／c','Δ pp','95% 條件區間','Holm p','原規則狀態'],rows,
 '每次家族為六。區間未作多重比較調整；狀態沿用當次預註冊，RULED_OUT 指 +10 pp 目標，並非零效應或等價。來源 [E2、E4、E12]。',[1.3,1.4,1.4,1.6,4.1,2,4.4])

# CONFORM - OFF is an out-of-family exploratory contrast, so it is not part of the
# preregistered primary recomputation stored in verified_evidence.json. The values
# below were recomputed from the same archived rows and cross-checked cell by cell
# against examples/verdicts.py (harness.gain_comes_from_executable_gate_and_resample).
CO={'LCB3_medium':('14/4','+7.41'),'LCB3_hard':('3/0','+5.56'),
    'HumanEvalPlus':('21/3','+11.54'),'MBPPPlus':('26/8','+4.85'),'pooled':('64/15','+6.84')}
rows=[]
for key,label in [('LCB3_medium','LCB v3 medium'),('LCB3_hard','LCB v3 hard'),('HumanEvalPlus','HumanEval+'),('MBPPPlus','MBPP+')]:
    v=DATA['cross_sets'][key];p=v['pairs']['HMIX_vs_CONFORM'];lo,hi=p['conditional_ci_pp']
    rows.append([label,str(p['n'])]+[str(v['arms'][a]['correct']) for a in ['OFF','CONFORM','HMIX']]
                +[f"{p['b']}/{p['c']}\n{p['delta_pp']:+.2f}",f'[{lo:.2f}, {hi:.2f}]','\n'.join(CO[key])])
v=DATA['cross_primary'];p=v['pairs']['HMIX_vs_CONFORM'];lo,hi=p['conditional_ci_pp']
rows.append(['合併',str(p['n'])]+[str(v['arms'][a]['correct']) for a in ['OFF','CONFORM','HMIX']]
            +[f"{p['b']}/{p['c']}\n{p['delta_pp']:+.2f}",f'[{lo:.2f}, {hi:.2f}]','\n'.join(CO['pooled'])])
table('cross','表 6　跨題庫交付件數與主要配對比較',
 ['題目集','n','OFF','CONFORM','H-MIX','H−C b／c 與 Δ pp','H−C 95% 條件區間','C−O b／c 與 Δ pp'],rows,
 'OFF、CONFORM、H-MIX 欄為成功件數，分母為同列 n。H−C 為 H-MIX 減 CONFORM，C−O 為 CONFORM 減 OFF；後者不在任何預註冊家族內，p 值未校正。區間未作多重比較調整。逐集絕對比例是兩種推論條件的混合，各集的兩台後端呼叫配比不同（見表 B2），不可單獨引用，也不可用於集間難度比較或模型排名；可比的是塊內配對的 b 與 c。主要 Holm p 為 H−C .3409、H−O 2.02 × 10⁻⁸，家族為二。來源 [E6、E12]。',[2.7,0.8,1.0,1.7,1.3,2.0,2.5,2.0])

table('hypotheses','表 7　假設、階段、統計量與事前規則判定',
 ['假設','對應階段','統計量與分母','事前規則判定'],[
 ['H1　修訂流程相對單次生成提高交付成功率','R460；R460R 五次',
  'H-PI、H-OC、H-MIX 對 OFF 共十五格，Δ 介於 +17.5 至 +29.2 pp；每次 n = 120，第五次為 118 至 119',
  '十五格全部通過各次六項家族的 Holm 校正；規則成立'],
 ['H2　H-MIX 相對 CONFORM 提高交付成功率（主假設）','R460；R460R 五次；R529',
  'R460 +13.33 pp（22/6、n = 120）；五次 +5.83／+4.17／+0.83／+2.50／+4.31 pp；跨題庫合併 +1.12 pp（31/23、n = 716）',
  'R460 列 EFFECTIVE；五次 0/5 通過 Holm，依規則逐次照實列；R529 p_adj .3409，列 INCONCLUSIVE'],
 ['H3　固定候選後驗收選擇相對行為多數決','r446、r448、r449b、r449c（EQ5）；R460R 為家族外探索比較',
  'EQ5 四階段 +4.04／+3.50／+8.33／+4.23 pp，未校正 p .0135／.0241／.0414／.0963；R460R 的 OFF5 對 CONFORM 五次 −10.83／−10.83／−15.83／−4.17／−1.74 pp',
  'EQ5 未設合併家族，逐階段照實列；R460R 五次同號、三次未校正 p < .05，判為同號未解析'],
 ['H4　假交付方向（依階段區分）','R460 的 P-H4；R460R 的 P-R4；R529 的 P-X4',
  'R460 H-MIX 14/120 對 CONFORM 29/120；R460R 五次皆較低；R529 四集為 6/135 對 9/135、7/54 對 11/54、6/156 對 6/156、58/371 對 56/371',
  '各階段均以方向性預測記分，未另配置確認性檢定；R460 未命中、R460R 命中、R529 未命中']],
 '每一列的判定依當階段的預註冊規則作成，同名標籤跨階段不可互引。H3 在 R460R 的比較與 H4 的三個階段皆不屬於任何預註冊 Holm 家族，其 p 值未經校正。來源 [E1–E6、E8]。',[3.3,2.9,7.0,4.4])

table('backends','表 B2　R529 逐後端的交付件數與推論模式',
 ['題目集','後端推論模式','OFF','CONFORM','H-MIX'],[
 ['LCB v3 medium','啟用 reasoning','69/75','74/75','74/75'],
 ['LCB v3 medium','未啟用','46/60','51/60','52/60'],
 ['LCB v3 hard','啟用 reasoning','25/34','26/34','28/34'],
 ['LCB v3 hard','未啟用','13/20','15/20','15/20'],
 ['HumanEval+','啟用 reasoning','87/100','94/100','94/100'],
 ['HumanEval+','未啟用','42/56','53/56','54/56'],
 ['MBPP+','啟用 reasoning','174/231','184/231','187/231'],
 ['MBPP+','未啟用','103/140','111/140','112/140']],
 '每格為成功交付件數與該後端在該集所承接的題數。塊隨機落在兩台，塊內同題三臂固定同一台，因此配對算術不受兩台差異影響。本表為描述性，不構成兩台之間的效能比較，也不用於任何仲裁。來源 [E6]。',[3.2,3.0,2.4,2.4,2.4])

table('predictions','表 E1　各階段事前預測逐條對帳',
 ['階段','編號','事前窗','實測','判定'],[
 ['R445','P-E1 至 P-E8','八項事前預測，含交付率窗與區間半寬','七項落在窗內；未命中者為 P-E2，併庫條件區間半寬 3.29 pp 大於預測的 3.0 pp','七命中、一未命中'],
 ['R460','P-H0','前六十題 OFF 交付率落在 38.3% 至 68.3%','60.0%','命中'],
 ['R460','P-H1','三個 H 臂皆高於 OFF','+17.50／+21.67／+25.83 pp','命中'],
 ['R460','P-H2','至少一個 H 臂高於 CONFORM','三臂差值皆為正','命中'],
 ['R460','P-H3','呼叫／題落在 1.8 至 3.2','1.51／2.36／1.36','未命中'],
 ['R460','P-H4','假交付高於 CONFORM 且不超過 35%','14.2%／15.0%／11.7%，低於 CONFORM 的 24.2%','未命中（方向相反）'],
 ['R460','P-H5','H-OC 與 H-MIX 的 loader 結束輪次不超過 0.5%','1.06% 與 6.13%','未命中'],
 ['R460','P-H6','H-MIX token 除以 H-PI 不超過 0.8','0.60','命中'],
 ['R460','P-H7','撞牆鐘比例不超過 5%','0／0／0','命中'],
 ['R460','P-H8','無程式碼輪次不超過 3%','0／0／0','命中'],
 ['R460','P-H9','兩條都顯著贏 H-PI 才追加第四臂','H-OC 對 H-PI p = .40；H-MIX 對 H-PI p = .041','未觸發，不列入計分'],
 ['R460R','P-R1','五次 Δ 皆大於零','5/5','命中 5/5'],
 ['R460R','P-R2','五次 Δ 皆不小於 +10 pp','0/5','未命中 5/5'],
 ['R460R','P-R3','五次皆通過 Holm 校正','0/5','未命中 5/5'],
 ['R460R','P-R4','五次假交付 H-MIX 低於 CONFORM','5/5','命中 5/5'],
 ['R460R','P-R5','五次 token／題不超過 CONFORM 的 1.2 倍','3/5','三次命中、兩次未命中'],
 ['R529','P-X1','四集 H-MIX 對 CONFORM 差值皆大於零','+0.74／+3.70／+0.64／+1.08 pp','命中'],
 ['R529','P-X2','四集 H-MIX 對 OFF 差值皆大於零','+8.15／+9.26／+12.18／+5.93 pp','命中'],
 ['R529','P-X3','MBPP+ 與 HumanEval+ 的差值皆小於 LCB v3 兩層的最小值 0.74','HumanEval+ 0.64 小於 0.74；MBPP+ 1.08 大於 0.74','未命中'],
 ['R529','P-X4','四集假交付皆為 H-MIX 低於 CONFORM','LCB 兩層成立（6 對 9、7 對 11）；HumanEval+ 6 對 6；MBPP+ 58 對 56','未命中'],
 ['R529','P-X5','四集 token／題皆不超過 CONFORM 的 1.2 倍','1.37／1.07／1.69／1.62 倍','未命中'],
 ['R529','P-X6','hard 層差值不小於 medium 層','3.70 不小於 0.74，但 hard 的不一致對僅 8','命中'],
 ['R529','P-X7','四集呼叫／題皆落在 1.2 至 2.5','1.14／1.24／1.09／1.11','未命中'],
 ['R529','P-X8','每塊 infra_void 不超過 5%','全部為 0','命中']],
 'R460 共九項列入計分，六項命中、三項未命中；P-H9 是階段觸發條件，不計分。未命中的窗一律保留原值，不回改。R445 的八項僅摘錄未命中者的數值，逐條原文見該階段收官檔。來源 [E1–E6、E8]。',[1.5,2.0,6.0,6.5,2.6])

table('academic','表 8　研究程序與尚未完成的要求',
 ['程序','已完成與證據','哪裡未做到'],[
 ['事前假設','主要門檻、家族、停止規則保存於版本紀錄 [E1、E3、E5]','R440P 為探索；H4 方向分階段不同'],
 ['事前分析','多數判準已有機器可讀欄位及檢查','R529 分析器晚於首塊完成；未讀資料不可獨立證實'],
 ['效果與不確定性','主表列分母、配對、效果、區間與 Holm p [E12]','條件區間未校正；不能代表全研究家族控制'],
 ['複製','同題庫五次新種子；跨題庫四集 [E4、E6]','重用題目；單模型；後端負載與模式不同'],
 ['重算','歷史三方核對；本文另一路徑重算 [E12]','同專案資料核對，不是外部獨立複製'],
 ['量具驗證','人工故障注入、回歸及鏈驗簽 [E9–E10]','未驗證自然洩漏偵測率；掃描有跳過項'],
 ['資料可得性','保存結果列、呼叫、收據及索引 [E11]','部分題庫需另取得；未全符合統一證據包格式'],
 ['偏離與負向結果','保留作廢塊、事故及主張降級 [E1–E6]','缺少外部事前評審與第三方不可變註冊']],
 '預註冊與透明報告支援可查核性，不能替代實驗效度或外部審查 [16]。',[2.4,7.2,6.6])

table('incidents','表 B1　研究事故的處置與剩餘影響',
 ['事件','已採取處置','推論影響'],[
 ['R440P 重放誤用隱藏標籤','改為僅用可見資訊選擇；保留更正 [E7]','屬失效分析版本，不能作有效交付證據'],
 ['R460 後端兩次崩潰','作廢塊隔離；改端點完成剩餘塊 [E1–E2]','拓撲在部分結果存在後修訂；執行時間不同'],
 ['V/GT v1 九十筆偽陽性','依來源建立 v2；配合已知故障注入 [E2、E4]','降低噪音但新增豁免；零命中依賴覆蓋範圍'],
 ['R460R 共租與併發假警報','揭露共租；以完成時刻減延遲重建窗口 [E4]','仲裁值未變；負載差異仍存在'],
 ['排程器中文字串解碼失敗','修正解碼及截斷處理 [E6]','造成排程停頓；非模型錯誤'],
 ['r5 模型崩潰與自動卸載','修正重載策略；七筆作廢另列 [E4、E6]','不同臂分母；需要缺失敏感度分析'],
 ['R529 後端 reasoning 不一致','補探針與描述欄；後續固定設定 [E6]','歷史資料仍為混合模式，不能追溯同質化']],
 '處置修復執行或量具，不使已存在的設計限制消失。',[3.5,6.3,6.4])

table('claims','表 C1　文獻與實驗各自能支持的主張',
 ['主張','可用依據','本文採取的界線'],[
 ['回饋修訂有研究基礎','Self-Debugging、Reflexion [6–7]','承接既有工作；不主張首創'],
 ['修復不一定優於重抽','Olausson [8]；R460R、R529','外部研究支持問題的重要性；本文效果靠自己的配對資料'],
 ['多代理可能一起出錯','Kim [9]','不將相關錯誤改寫成所有投票必然無效'],
 ['驗收可改善候選選擇','EQ5；CONFORM 對照 [E8]','同候選局部證據；不外推所有測資與模型'],
 ['H-MIX 額外收益已穩定','現有複製未達事前規則 [E4、E6]','不採用此主張；保留差值、區間及未顯著結果'],
 ['簽章使紀錄可核驗','PeerReview、CT 原則；逐筆驗簽 [13–14、E12]','限定鏈內一致性；不等於內容真實或完整歷史'],
 ['系統提高人類信任','需要構念與使用者測量 [15]','本文未測量，不作成效主張'],
 ['研究遵守全部學術要求','表 8 的已完成與缺口','不以內部查核冒充外部同行評審']],
 '書目以原始研究、官方論文頁或作者稿為主。文獻整理屬主題式回顧，未執行系統性綜述的窮盡檢索。',[4.1,5.5,6.6])

EVIDENCE=[
 ('E1','R460 事前設計與拓撲修訂',['DECISION_20260907_R460_HARNESS_PREREG.md']),
 ('E2','R460 初次結果與更正',['DECISION_20260911_R460_FABLE_AUDIT_HARNESS.md']),
 ('E3','五次複製的事前規則',['DECISION_20260911_R460R_FIVE_REPLICATIONS_PREREG.md']),
 ('E4','五次複製完整收官，以第八節補記為準',['DECISION_20260912_R460R_FABLE_AUDIT_REPLICATIONS.md']),
 ('E5','R529 跨題庫事前規則',['DECISION_20260911_R529_CROSS_BANK_PREREG.md']),
 ('E6','R529 收官與後端模式更正，含第十一至十三節',['DECISION_20260912_R529_FABLE_AUDIT_CROSS_BANK.md']),
 ('E7','候選池的探索性重放',['DECISION_20260903_R440P_CONFORMANCE_GATE.md']),
 ('E8','CONFORM 與 EQ5 階段結果',['CONCLUSION_20260904_R445_CONFORM_SETTLEMENT.md','CONCLUSION_20260904_R446_EQUAL_BUDGET.md','DECISION_20260907_R449C_FABLE_AUDIT_UNRESOLVED.md','docs/VACANT_COMPLETE_2026-09-12.md']),
 ('E9','多方執行與攻擊邊界',['DECISION_20260906_R453_FABLE_AUDIT_REAL_MULTIPARTY.md','DECISION_20260906_R454_FABLE_AUDIT_NAMED_DISSENT.md','DECISION_20260905_R449_PEEREXEC_ARCHITECTURE_AUDIT.md']),
 ('E10','簽章格式與原始驗簽報告',['vacant/logbook.py','vacant/canonical.py','ops/gain/replay/r460r/receipts_verify.json','ops/gain/replay/r460r/receipts_verify_r4r5.json','ops/gain/replay/r529/receipts_verify.json']),
 ('E11','題庫、歷程與實作索引',['runs/INDEX.md','vacant/codebench.py','docs/JOURNEY_2026-09-13.md','docs/HMIX_ARCHITECTURE_2026-09-11.md']),
 ('E12','本文從結果列及公開簽章重新核對的程式與輸出',['docs/paper_2026-09-14/verify_evidence.py','docs/paper_2026-09-14/verified_evidence.json']),
]
for _,_,paths in EVIDENCE:
    for p in paths: assert (ROOT/p).exists(),p

def figure_pdf(name,w,h,draw):
    if (TMP/(name+'.png')).exists():return
    pdf=TMP/(name+'.pdf');c=canvas.Canvas(str(pdf),pagesize=(w,h));draw(c);c.save()
    img=convert_from_path(str(pdf),dpi=180,poppler_path='/Users/cosmopig/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override')[0]
    img.save(TMP/(name+'.png'))

def draw_arch(c):
    ink=HexColor('#20262C');accent=HexColor('#315D71')
    def box(x,y,w,h,title,sub):
        c.setFillColor(HexColor('#F5F6F7'));c.setStrokeColor(ink);c.roundRect(x,y,w,h,5,fill=1)
        c.setFillColor(ink);c.setFont('Helvetica-Bold',11);c.drawCentredString(x+w/2,y+h-19,title)
        c.setFont('Helvetica',9);c.drawCentredString(x+w/2,y+14,sub)
    def arrow(x,y,xx,yy):
        c.setStrokeColor(accent);c.setLineWidth(1.1);c.line(x,y,xx,yy)
        import math
        a=math.atan2(yy-y,xx-x)
        for da in [-.5,.5]:c.line(xx,yy,xx-6*math.cos(a+da),yy-6*math.sin(a+da))
    box(12,180,150,55,'Requirement + V','Visible tests / contract')
    box(210,180,150,55,'Generate / revise','OFF, CONFORM, H arms')
    box(405,180,150,55,'Run visible tests','Sandboxed acceptance')
    arrow(162,207,210,207);arrow(360,207,405,207)
    box(405,78,150,55,'Deliver or refuse','Gate and budget rules')
    arrow(480,180,480,133)
    c.setFont('Helvetica',9);c.setFillColor(ink);c.drawString(487,153,'PASS / stop')
    c.setStrokeColor(accent);c.line(435,180,435,151);c.line(435,151,285,151);arrow(285,151,285,180)
    c.setFont('Helvetica',9);c.drawString(300,138,'FAIL: visible feedback only')
    box(12,78,150,55,'Hidden evaluation GT','After selection; no feedback')
    c.setDash(4,3);arrow(405,95,162,95);c.setDash()
    c.setFont('Helvetica',9);c.drawCentredString(285,104,'Preserved final candidate')
    box(170,0,230,53,'Signed receipt chain','Attempts + verdicts of instrumented arms')
    arrow(480,78,400,27)
    c.setFont('Helvetica-Oblique',8);c.drawString(15,60,'GT is a finite test oracle, not complete truth.')
figure_pdf('architecture',570,250,draw_arch)

def draw_forest(c):
    left,right=100,440;bottom,top=54,285
    def x(v):return left+(v+10)/32*(right-left)
    c.setFont('Helvetica-Bold',13);c.drawString(15,323,'H-MIX minus CONFORM')
    c.setFont('Helvetica',9);c.drawString(15,307,'Paired differences and unadjusted 95% conditional intervals')
    c.setStrokeColor(HexColor('#BBBBBB'));c.setDash(3,3);c.line(x(0),bottom,x(0),top+6);c.setDash()
    c.setStrokeColor(HexColor('#AAAAAA'));c.line(left,bottom,right,bottom)
    for tick in [-10,-5,0,5,10,15,20]:
        c.line(x(tick),bottom-4,x(tick),bottom)
        c.setFillColor(HexColor('#333333'));c.setFont('Helvetica',9);c.drawCentredString(x(tick),bottom-17,str(tick))
    for i,k in enumerate(labels):
        y=top-i*39;v=DATA[k]['pairs']['HMIX_vs_CONFORM'];lo,hi=v['conditional_ci_pp'];col=HexColor('#223A4A') if i==0 else HexColor('#54798B')
        c.setFillColor(HexColor('#222222'));c.setFont('Helvetica-Bold' if i==0 else 'Helvetica',11);c.drawString(20,y-4,k)
        c.setStrokeColor(col);c.setFillColor(col);c.setLineWidth(1.5);c.line(x(lo),y,x(hi),y)
        c.line(x(lo),y-4,x(lo),y+4);c.line(x(hi),y-4,x(hi),y+4);c.circle(x(v['delta_pp']),y,3.2,fill=1,stroke=0)
        c.setFillColor(HexColor('#222222'));c.setFont('Helvetica',10);c.drawString(455,y-4,f"{v['delta_pp']:+.2f} pp   p_adj={v['p_holm']:.4f}")
    c.setFont('Helvetica',10);c.drawCentredString((left+right)/2,16,'Difference in percentage points')
figure_pdf('forest',630,344,draw_forest)

FIGURES={
 'architecture':('圖 1　可見驗收 隱藏評分與收據紀錄的資料流', 'GT 僅對保存的最終候選事後評分，不回流生成或出貨選擇。收據只涵蓋有加入記錄功能的實驗臂。'),
 'forest':('圖 2　初次實驗與五次複製的主要配對效果', 'r1 至 r5 皆重用相同 120 題；r5 主比較共同分母為 116。每次各自進行六項 Holm 校正；圖中區間未作多重比較調整。來源 [E2、E4、E12]。')}

doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.25);sec.bottom_margin=Cm(2.15);sec.left_margin=Cm(2.25);sec.right_margin=Cm(2.25)
sec.footer_distance=Cm(.95)
def font_style(style,size,bold=False,face='Times New Roman',east='Songti TC'):
    s=doc.styles[style];s.font.name=face;s.font.size=Pt(size);s.font.bold=bold;s.font.color.rgb=RGBColor(0,0,0)
    rf=s.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in list(rf.attrib):
        if 'Theme' in attr: del rf.attrib[attr]
    rf.set(qn('w:eastAsia'),east)
    return s
s=font_style('Normal',11.5);s.paragraph_format.line_spacing=1.42;s.paragraph_format.space_after=Pt(6)
s.paragraph_format.widow_control=True
font_style('Title',22,True,east='Heiti TC').paragraph_format.space_after=Pt(12)
for name,size in [('Heading 1',16),('Heading 2',13)]:
    s=font_style(name,size,True,east='Heiti TC');s.paragraph_format.space_before=Pt(15);s.paragraph_format.space_after=Pt(8);s.paragraph_format.keep_with_next=True
s=font_style('Caption',9.5);s.paragraph_format.space_after=Pt(5)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
doc.core_properties.title='可執行驗收支援的 LLM 程式代理交付'
doc.core_properties.subject='Vacant 的機制比較、複製實驗與可究責紀錄'
doc.core_properties.author='';doc.core_properties.keywords='LLM; executable acceptance; accountability; replication'
# The bundled default template contains title borders; remove all style borders.
for el in list(doc.styles.element.iter(qn('w:pBdr'))):el.getparent().remove(el)

def add_link(p,label,url):
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'254E69');pr.append(color);r.append(pr)
    t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

def rich(p,text):
    # Latin super/subscript font fallback is more robust as Unicode in the renderer.
    for part in re.split(r'(https?://[^\s]+)',text):
        if part.startswith('http'):add_link(p,part,part)
        else:p.add_run(part)

def note(text):
    p=doc.add_paragraph(style='Caption');rich(p,text);p.paragraph_format.keep_with_next=False
    return p

def add_table(key):
    title,head,rows,foot,widths=TABLES[key]
    p=doc.add_paragraph(title,style='Caption');p.paragraph_format.keep_with_next=True
    p.runs[0].bold=True
    t=doc.add_table(rows=1,cols=len(head));t.autofit=False
    if widths:
        total=sum(widths);widths=[w/total*16.5 for w in widths]
        for col,w in zip(t.columns,widths):col.width=Cm(w)
    for i,h in enumerate(head):t.rows[0].cells[i].text=h
    for row in rows:
        cells=t.add_row().cells
        for i,txt in enumerate(row):cells[i].text=str(txt)
    for ridx,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();avoid=OxmlElement('w:cantSplit');trpr.append(avoid)
        if ridx==0:
            rep=OxmlElement('w:tblHeader');trpr.append(rep)
        for i,cell in enumerate(row.cells):
            if widths:cell.width=Cm(widths[i])
            tc=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
            for tag,val in [('top','80'),('bottom','80'),('left','75'),('right','75')]:
                node=OxmlElement('w:'+tag);node.set(qn('w:w'),val);node.set(qn('w:type'),'dxa');marg.append(node)
            tc.append(marg)
            if ridx==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F0F1F2');tc.append(sh)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.16
                p.paragraph_format.keep_with_next=False
                for r in p.runs:r.font.size=Pt(9.3);r.bold=(ridx==0)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','bottom','insideH']:
        el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'B9BDC1');borders.append(el)
    t._tbl.tblPr.append(borders)
    note(foot)

expanded=[];template=(BASE/'manuscript_template.md').read_text()
replacements={'验收':'驗收','題数':'題數','来源':'來源','–':'-','後续':'後續','この':'此'}
for a,b in replacements.items():template=template.replace(a,b)
for key,values in list(TABLES.items()):
    title,head,rows,foot,widths=values
    def clean(x):
        for a,b in replacements.items():x=x.replace(a,b)
        return x
    TABLES[key]=(clean(title),list(map(clean,head)),[[clean(str(x)) for x in r] for r in rows],clean(foot),widths)

in_refs=False;is_abs=False
for block in template.split('\n\n'):
    block=block.strip()
    if not block:continue
    if block.startswith('# 可執行'):
        for line in block.splitlines():
            if line.startswith('# '):doc.add_paragraph(line[2:],'Title');expanded.append(line)
            elif line.startswith('##TITLE_SUB '):
                p=doc.add_paragraph(line[12:]);p.runs[0].font.size=Pt(15);p.runs[0].bold=True;expanded.append(line[12:])
            elif line.startswith('##TITLE_EN '):
                p=doc.add_paragraph(line[11:]);p.paragraph_format.space_after=Pt(3);p.runs[0].font.size=Pt(11);expanded.append(line[11:])
            elif line.startswith('##META '):
                p=note(line[7:]);p.paragraph_format.space_before=Pt(10);expanded.append(line[7:])
        continue
    if block.startswith('##PAGE'):
        doc.add_page_break();block=block[6:].strip()
        if not block:continue
    if block.startswith('[[TABLE:'):
        key=block[8:-2];add_table(key);title,head,rows,foot,_=TABLES[key]
        expanded.append(title+'\n\n|'+'|'.join(head)+'|\n|'+'|'.join(['---']*len(head))+'|\n'+'\n'.join('|'+'|'.join(x.replace('\n','<br>') for x in row)+'|' for row in rows)+'\n\n'+foot);continue
    if block.startswith('[[FIGURE:'):
        key=block[9:-2];title,foot=FIGURES[key]
        p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(TMP/(key+'.png')),width=Cm(16.2))
        note(title+'。'+foot);expanded.append(title+'\n\n'+foot);continue
    if block=='[[EVIDENCE_INDEX]]':
        for eid,label,paths in EVIDENCE:
            p=doc.add_paragraph(f'[{eid}] {label}');p.paragraph_format.keep_with_next=True;p.runs[0].bold=True
            expanded.append(f'[{eid}] {label}')
            for path in paths:
                p=doc.add_paragraph();p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.1
                add_link(p,path,'https://github.com/cosmopig/Vacant/blob/44be37fe52bac148ecd2835aa384c804c612ffea/'+path) if eid!='E12' else p.add_run(path)
                # Allow long file identifiers to wrap at underscore boundaries.
                for text in p._p.iter(qn('w:t')):
                    text.text=(text.text or '').replace('_','_\u200b').replace('/','/\u200b')
                for r in p.runs:r.font.size=Pt(9)
                expanded.append(path)
        continue
    if block.startswith('### '):
        doc.add_paragraph(block[4:],'Heading 2');expanded.append(block);continue
    if block.startswith('## '):
        title=block[3:];doc.add_paragraph(title,'Heading 1');in_refs=(title=='參考文獻');is_abs=title in ['摘要','Abstract'];expanded.append(block);continue
    p=doc.add_paragraph();rich(p,block)
    if in_refs:
        p.paragraph_format.left_indent=Cm(.75);p.paragraph_format.first_line_indent=Cm(-.75);p.paragraph_format.line_spacing=1.18;p.paragraph_format.space_after=Pt(8)
        for r in p.runs:r.font.size=Pt(10)
    elif not is_abs and not block.startswith('關鍵詞'):
        p.paragraph_format.first_line_indent=Cm(.65)
    expanded.append(block)

doc.save(BASE/(STEM+'.docx'))
(BASE/'manuscript.md').write_text('\n\n'.join(expanded)+'\n')
# Reference list lives in manuscript_template.md and is harvested here.
# Maintainer note for [7] Reflexion: the URL and the author list must move together.
# The NeurIPS 36 conference version has five authors (Shinn, Cassano, Gopinath,
# Narasimhan, Yao); the arXiv preprint (2303.11366) has six, adding E. Berman.
# Swapping the link to arXiv therefore requires adding Berman, and keeping the
# conference DOI/page range requires leaving him out.
refs=[]
for block in template.split('\n\n'):
    if re.match(r'^\[\d+\]',block):
        refs.append({'id':int(re.match(r'^\[(\d+)\]',block)[1]),'citation':block,'urls':re.findall(r'https?://\S+',block),'checked_on':'2026-09-14'})
source_root=Path('/Users/cosmopig/Library/Mobile Documents/com~apple~CloudDocs/專題/參考文獻')
local=[]
for pat in ['原文/*.pdf','2026-08-06_agent信任/pdf/*Kim*','2026-08-06_信任定義/pdf/*Koerber*']:
    for path in source_root.glob(pat):
        local.append({'path':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'used':not path.name.startswith('Final-Web')})
(BASE/'source_manifest.json').write_text(json.dumps({'references':refs,'local_originals':local,'internal':EVIDENCE,'baseline_commit':'44be37fe52bac148ecd2835aa384c804c612ffea'},ensure_ascii=False,indent=2)+'\n')
print('DOCX',BASE/(STEM+'.docx'))
print('References',len(refs),'Tables',len(TABLES),'CJK characters',len(re.findall(r'[\u4e00-\u9fff]','\n'.join(expanded))))
