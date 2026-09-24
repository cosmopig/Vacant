import regex as re
import glob
import os
from openpyxl import load_workbook

def task_func(directory_path='./xlsx_files/'):
    count = 0
    # Get all .xlsx files in the directory
    files = glob.glob(os.path.join(directory_path, "*.xlsx"))
    
    for file_path in files:
        try:
            wb = load_workbook(file_path)
            modified = False
            for sheet in wb.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value and isinstance(cell.value, str):
                            # Check if the string contains double quotes
                            if '"' in cell.value:
                                # Replace " with \"
                                new_value = cell.value.replace('"', '\\"')
                                cell.value = new_value
                                modified = True
            
            if modified:
                wb.save(file_path)
            count += 1
        except Exception:
            # If a file cannot be processed, we skip it or handle as needed.
            # The goal says "Protects all double quotes in all Excel (.xlsx) files"
            # Usually this implies processing them if possible.
            pass
            
    return count
