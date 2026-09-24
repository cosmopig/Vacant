import os
from pathlib import Path
import pandas as pd
import docx

def task_func(source_directory: str, target_directory: str) -> int:
    source_path = Path(source_directory)
    if not source_path.exists() or not source_path.is_dir():
        raise FileNotFoundError(f"Source directory {source_directory} does not exist.")

    target_path = Path(target_directory)
    target_path.mkdir(parents=True, exist_ok=True)

    count = 0
    for file_path in source_path.iterdir():
        if file_path.is_file():
            ext = file_path.suffix.lower()
            if ext in ['.txt', '.docx', '.xlsx', '.csv']:
                try:
                    content = ""
                    if ext == '.txt':
                        content = file_path.read_text(encoding='utf-8')
                    elif ext == '.docx':
                        doc = docx.Document(file_path)
                        texts = []
                        for para in doc.paragraphs:
                            if para.text.strip():
                                texts.append(para.text)
                        for table in doc.tables:
                            for row in table.rows:
                                for cell in row.cells:
                                    if cell.text.strip():
                                        texts.append(cell.text)
                        content = "\n".join(texts)
                    elif ext == '.xlsx':
                        xl = pd.ExcelFile(file_path)
                        all_values = []
                        for sheet in xl.sheet_names:
                            df = xl.parse(sheet)
                            # Convert all values to string and flatten, filtering out NaNs
                            vals = df.astype(str).values.flatten()
                            all_values.extend([v for v in vals if v != 'nan' and v != 'None'])
                        content = "\n".join(all_values)
                    elif ext == '.csv':
                        df = pd.read_csv(file_path)
                        vals = df.astype(str).values.flatten()
                        content = "\n".join([v for v in vals if v != 'nan' and v != 'None'])

                    output_filename = file_path.stem + ".csv"
                    output_file_path = target_path / output_filename
                    
                    df_out = pd.DataFrame({'Text': [content]})
                    df_out.to_csv(output_file_path, index=False, header=True)
                    count += 1
                except Exception:
                    continue
    return count
