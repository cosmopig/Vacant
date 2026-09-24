import os
from pathlib import Path
import pandas as pd
import docx

def task_func(source_directory: str, target_directory: str) -> int:
    source_path = Path(source_directory)
    if not source_path.exists() or not source_path.is_dir():
        raise FileNotFoundError(f"Source directory {source_directory} does not exist.")

    target_path = Path(target_directory)
    target_path.mkdir(parents=True, exist_ok=True)

    success_count = 0
    
    # Iterate over files in the source directory.
    for file_path in source_path.iterdir():
        if file_path.is_file():
            ext = file_path.suffix.lower()
            if ext in ['.txt', '.docx', '.xlsx', '.csv']:
                # The target filename should have the same base name but .csv extension.
                base_name = file_path.stem
                target_file = target_path / f"{base_name}.csv"
                
                try:
                    content = ""
                    if ext == '.txt':
                        content = file_path.read_text(encoding='utf-8')
                    elif ext == '.docx':
                        doc = docx.Document(str(file_path))
                        full_text = []
                        for para in doc.paragraphs:
                            full_text.append(para.text)
                        for table in doc.tables:
                            for row in table.rows:
                                for cell in row.cells:
                                    full_text.append(cell.text)
                        content = "\n".join(full_text)
                    elif ext == '.xlsx':
                        df = pd.read_excel(str(file_path))
                        content = df.to_string(index=False)
                    elif ext == '.csv':
                        # Read the CSV and convert its content to a string.
                        df = pd.read_csv(str(file_path))
                        content = df.to_string(index=False)
                    
                    # Create a DataFrame with a single 'Text' column.
                    df_out = pd.DataFrame({'Text': [content]})
                    df_out.to_csv(target_file, index=False)
                    success_count += 1
                except Exception:
                    # If conversion fails for some reason (e.g. corrupted file), skip it.
                    continue
                    
    return success_count
